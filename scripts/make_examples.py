"""Génère les fichiers de démonstration SYNTHÉTIQUES livrés avec l'application (aucune donnée réelle).

Usage : python scripts/make_examples.py
Les sorties sont écrites dans backend/kpibridge/examples/ et sont déterministes (graine fixe).
"""

from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "backend" / "kpibridge" / "examples"
FONT = ROOT / "backend" / "kpibridge" / "exports" / "fonts" / "NotoSans-Regular.ttf"

PRODUCTS = [
    ("Stylo bille", "Papeterie", Decimal("1.90")),
    ("Cahier A4", "Papeterie", Decimal("3.50")),
    ("Agenda 2026", "Papeterie", Decimal("12.00")),
    ("Clé USB 64 Go", "Informatique", Decimal("14.90")),
    ("Souris sans fil", "Informatique", Decimal("19.99")),
    ("Casque audio", "Informatique", Decimal("49.00")),
    ("Lampe de bureau", "Mobilier", Decimal("34.50")),
    ("Chaise ergonomique", "Mobilier", Decimal("189.00")),
]
REGIONS = ["Nord", "Sud", "Est", "Ouest"]


def sales_rows(seed: int = 42, n_orders: int = 260) -> list[dict]:
    rnd = random.Random(seed)
    rows = []
    start = date(2025, 1, 1)
    for i in range(1, n_orders + 1):
        d = start + timedelta(days=rnd.randrange(0, 365))
        region = rnd.choice(REGIONS)
        client = f"C{rnd.randrange(1, 90):03d}"
        for _ in range(rnd.randint(1, 4)):
            name, cat, price = rnd.choice(PRODUCTS)
            qty = rnd.randint(1, 5)
            rows.append({
                "order_id": f"CMD-{i:05d}",
                "date": d,
                "client": client,
                "region": region,
                "product": name,
                "category": cat,
                "qty": qty,
                "unit_price": price,
                "amount": price * qty,
                "cost": (price * qty * Decimal("0.62")).quantize(Decimal("0.01")),
            })
    rows.sort(key=lambda r: (r["date"], r["order_id"]))
    return rows


def write_clean_csv(rows: list[dict]) -> None:
    with open(OUT / "ventes_demo.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["N° commande", "Date", "Client", "Région", "Produit", "Catégorie", "Quantité", "Prix unitaire", "Montant", "Coût", "Devise"])
        for r in rows:
            w.writerow([r["order_id"], r["date"].isoformat(), r["client"], r["region"], r["product"], r["category"], r["qty"], r["unit_price"], r["amount"], r["cost"], "EUR"])


EN = {
    "Stylo bille": "Ballpoint pen", "Cahier A4": "A4 notebook", "Agenda 2026": "2026 planner",
    "Clé USB 64 Go": "64 GB USB drive", "Souris sans fil": "Wireless mouse", "Casque audio": "Headphones",
    "Lampe de bureau": "Desk lamp", "Chaise ergonomique": "Ergonomic chair",
    "Papeterie": "Stationery", "Informatique": "IT accessories", "Mobilier": "Furniture",
    "Nord": "North", "Sud": "South", "Est": "East", "Ouest": "West",
}


def write_en_csv(rows: list[dict]) -> None:
    """Même jeu synthétique que ventes_demo.csv, avec en-têtes et libellés en anglais (démonstration en anglais)."""
    with open(OUT / "sales_demo_en.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Order ID", "Date", "Customer", "Region", "Product", "Category", "Quantity", "Unit price", "Amount", "Cost", "Currency"])
        for r in rows:
            w.writerow([r["order_id"].replace("CMD", "ORD"), r["date"].isoformat(), r["client"], EN[r["region"]], EN[r["product"]],
                        EN[r["category"]], r["qty"], r["unit_price"], r["amount"], r["cost"], "EUR"])


def fr_number(d: Decimal) -> str:
    s = f"{d:,.2f}"
    return s.replace(",", " ").replace(".", ",")


def write_fr_csv(rows: list[dict]) -> None:
    with open(OUT / "ventes_fr_point_virgule.csv", "w", newline="", encoding="utf-8-sig") as fh:
        fh.write("Export synthétique — ventes magasin (données fictives)\n\n")
        w = csv.writer(fh, delimiter=";")
        w.writerow(["Code client", "Date", "Produit", "Montant TTC", "Statut"])
        for i, r in enumerate(rows[:400]):
            status = "Annulée" if i % 37 == 0 else ("Remboursée" if i % 53 == 0 else "Payée")
            w.writerow([f"{int(r['client'][1:]):05d}", r["date"].strftime("%d/%m/%Y"), r["product"], fr_number(r["amount"]), status])
        total = sum((r["amount"] for r in rows[:400]), Decimal(0))
        w.writerow(["Total", "", "", fr_number(total), ""])


def write_xlsx(rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes par région"
    ws["A1"] = "Rapport régional (données fictives)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")
    ws.append([])
    ws.append(["Région", "Mois", "Chiffre d'affaires", "Commandes"])
    agg: dict[tuple[str, str], list] = {}
    for r in rows:
        key = (r["region"], r["date"].strftime("%Y-%m"))
        a = agg.setdefault(key, [Decimal(0), set()])
        a[0] += r["amount"]
        a[1].add(r["order_id"])
    for region in REGIONS:
        start = ws.max_row + 1
        for month in sorted({m for (rg, m) in agg if rg == region}):
            v = agg[(region, month)]
            ws.append([region, month, float(v[0]), len(v[1])])
        end = ws.max_row
        ws.append([f"Sous-total {region}", None, f"=SUM(C{start}:C{end})", f"=SUM(D{start}:D{end})"])
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Ce classeur est synthétique."
    wb.save(OUT / "classeur_regions_demo.xlsx")


def write_pdf(rows: list[dict]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(TTFont("NS", str(FONT)))
    doc = SimpleDocTemplate(str(OUT / "rapport_ventes_demo.pdf"), pagesize=A4)
    data = [["Commande", "Date", "Produit", "Montant (EUR)"]]
    for r in rows[:120]:
        data.append([r["order_id"], r["date"].isoformat(), r["product"], f"{r['amount']:.2f}"])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "NS"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.4, "#999999"),
                           ("BACKGROUND", (0, 0), (-1, 0), "#eeeeee")]))
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "NS"
    doc.build([Paragraph("Rapport de ventes — données synthétiques", styles["Title"]), Spacer(1, 12), t])


def table_image(rows: list[list[str]], path: Path, scale: int = 2, degrade: bool = False) -> None:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    font = ImageFont.truetype(str(FONT), 16 * scale)
    col_w = [150 * scale, 140 * scale, 150 * scale]
    h = 34 * scale
    img = Image.new("RGB", (sum(col_w) + 40 * scale, h * len(rows) + 40 * scale), "white")
    d = ImageDraw.Draw(img)
    y = 20 * scale
    for row in rows:
        x = 20 * scale
        for j, cell in enumerate(row):
            d.text((x + 6 * scale, y + 6 * scale), cell, fill="black", font=font)
            x += col_w[j]
        d.line([(20 * scale, y + h - 2), (x, y + h - 2)], fill="#999999", width=1)
        y += h
    if degrade:
        rnd = random.Random(1)
        img = img.resize((img.width // 3, img.height // 3)).filter(ImageFilter.GaussianBlur(1.2))
        px = img.load()
        for _ in range(img.width * img.height // 12):
            x, yy = rnd.randrange(img.width), rnd.randrange(img.height)
            px[x, yy] = (rnd.randrange(80, 200),) * 3
    img.save(path)


SCAN_ROWS = [["Mois", "Ventes", "Commandes"], ["Janvier", "12 450", "310"], ["Février", "11 980", "295"], ["Mars", "13 720", "342"], ["Avril", "14 105", "351"], ["Mai", "15 230", "377"], ["Juin", "14 870", "366"]]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = sales_rows()
    write_clean_csv(rows)
    write_en_csv(rows)
    write_fr_csv(rows)
    write_xlsx(rows)
    write_pdf(rows)
    table_image(SCAN_ROWS, OUT / "scan_tableau_demo.png")
    (OUT / ".README").write_text(
        "Fichiers de démonstration SYNTHÉTIQUES générés par scripts/make_examples.py (graine fixe).\n"
        "Aucune donnée réelle. / Synthetic demo files, no real data.\n",
        encoding="utf-8",
    )
    print("exemples écrits dans", OUT)


if __name__ == "__main__":
    main()
