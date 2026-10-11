"""Excel : XLSX (openpyxl, lecture seule, valeurs en cache uniquement) et XLS (xlrd).

Aucune formule n'est recalculée et aucun programme externe (Excel, LibreOffice) n'est lancé. Une formule sans valeur
en cache est marquée « indisponible ». Les macros ne sont jamais exécutées (les fichiers à macros sont refusés en amont).
"""

from __future__ import annotations

import datetime as dt
import posixpath
import zipfile
from pathlib import Path

import openpyxl
import xlrd
from defusedxml import ElementTree as DET
from openpyxl.utils import get_column_letter, range_boundaries

from ..errors import LimitExceeded, UserError
from .base import (
    FORMULA_NO_CACHE,
    IngestContext,
    RawTable,
    cached,
    clean_text,
    native_bool,
    native_date,
    native_number,
)
from .tables import structure_grid

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def _encode(value: object, number_format: str | None = None) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return native_bool(value)
    if isinstance(value, (int, float)):
        pct = bool(number_format and "%" in number_format)
        enc = native_number(value, percent=pct)
        return enc or None
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return native_date(value.date().isoformat())
        return native_date(value.isoformat(timespec="seconds"))
    if isinstance(value, dt.date):
        return native_date(value.isoformat())
    if isinstance(value, dt.time):
        return clean_text(value.isoformat())
    text = clean_text(str(value))
    return text or None


def _merged_ranges(path: Path) -> dict[str, list[str]]:
    """Lit les plages fusionnées de chaque feuille directement dans le XML (parseur sans entités externes)."""
    result: dict[str, list[str]] = {}
    with zipfile.ZipFile(path) as zf:
        wb = DET.fromstring(zf.read("xl/workbook.xml"))
        rels = DET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        for sheet in wb.findall("m:sheets/m:sheet", NS):
            name = sheet.get("name") or ""
            rid = sheet.get(f"{{{NS['r']}}}id")
            target = targets.get(rid)
            if not target:
                continue
            target = target.lstrip("/")
            full = target if target.startswith("xl/") else posixpath.normpath(posixpath.join("xl", target))
            if full not in zf.namelist():
                continue
            refs: list[str] = []
            with zf.open(full) as fh:
                for _event, elem in DET.iterparse(fh, events=("end",)):
                    if elem.tag == f"{{{NS['m']}}}mergeCell":
                        ref = elem.get("ref")
                        if ref:
                            refs.append(ref)
                    elif elem.tag == f"{{{NS['m']}}}row":
                        elem.clear()
            result[name] = refs
    return result


def _apply_merges(grid: list[list[str | None]], refs: list[str], table_notes: list[dict]) -> None:
    filled = 0
    for ref in refs:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
        except ValueError:
            continue
        r0, c0 = min_row - 1, min_col - 1
        if r0 >= len(grid) or c0 >= len(grid[r0]):
            continue
        value = grid[r0][c0]
        for r in range(r0, min(max_row, len(grid))):
            for c in range(c0, max_col):
                if (r, c) == (r0, c0):
                    continue
                while c >= len(grid[r]):
                    grid[r].append(None)
                if grid[r][c] in (None, ""):
                    grid[r][c] = value
                    filled += 1
    if filled:
        table_notes.append({"code": "merged_cells_filled", "params": {"count": filled, "ranges": len(refs)}})


def parse_xlsx(path: Path, name: str, ctx: IngestContext) -> list[RawTable]:
    limits = ctx.limits
    try:
        merges = _merged_ranges(path)
        wb_values = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        wb_formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (KeyError, zipfile.BadZipFile, OSError) as exc:
        raise UserError("corrupted_file") from exc
    selected = ctx.options.get("sheets")
    tables: list[RawTable] = []
    try:
        sheet_names = wb_values.sheetnames
        for s_idx, sheet_name in enumerate(sheet_names):
            if selected and sheet_name not in selected:
                continue
            ws_v = wb_values[sheet_name]
            ws_f = wb_formulas[sheet_name]
            grid: list[list[str | None]] = []
            srcs: list[str] = []
            formula_no_cache = 0
            formula_cached = 0
            cells = 0
            for r_idx, (row_v, row_f) in enumerate(zip(ws_v.iter_rows(), ws_f.iter_rows(), strict=False), start=1):
                if r_idx > limits.max_rows + 1000:
                    raise LimitExceeded("too_many_rows_sheet", sheet=sheet_name, limit=limits.max_rows)
                if len(row_v) > limits.max_cols + 50:
                    raise LimitExceeded("too_many_columns", limit=limits.max_cols, count=len(row_v))
                out: list[str | None] = []
                for cell_v, cell_f in zip(row_v, row_f, strict=False):
                    is_formula = getattr(cell_f, "data_type", None) == "f" or (
                        isinstance(getattr(cell_f, "value", None), str) and str(cell_f.value).startswith("=")
                    )
                    value = getattr(cell_v, "value", None)
                    fmt = getattr(cell_v, "number_format", None)
                    if is_formula:
                        if value is None:
                            out.append(FORMULA_NO_CACHE)
                            formula_no_cache += 1
                        else:
                            enc = _encode(value, fmt)
                            out.append(cached(enc) if enc else None)
                            formula_cached += 1
                    else:
                        out.append(_encode(value, fmt))
                cells += len(out)
                if cells > limits.max_cells:
                    raise LimitExceeded("too_many_cells", limit=limits.max_cells)
                grid.append(out)
                srcs.append(f"sheet={sheet_name};row={r_idx}")
            ctx.progress("extract", 0.2 + 0.6 * (s_idx + 1) / max(1, len(sheet_names)))
            if not grid:
                continue
            notes: list[dict] = []
            _apply_merges(grid, merges.get(sheet_name, []), notes)
            width = max(len(r) for r in grid)
            col_refs = [get_column_letter(i + 1) for i in range(width)]
            for t in structure_grid(grid, srcs, f"{name} › {sheet_name}", "sheet", col_refs=col_refs):
                t.notes.extend(notes)
                t.note("sheet", name=sheet_name)
                if formula_no_cache:
                    t.note("formula_without_cache", count=formula_no_cache)
                if formula_cached:
                    t.note("formula_cached_values", count=formula_cached)
                tables.append(t)
    finally:
        wb_values.close()
        wb_formulas.close()
    if not tables:
        raise UserError("no_data_found")
    return tables


def parse_xls(path: Path, name: str, ctx: IngestContext) -> list[RawTable]:
    limits = ctx.limits
    try:
        book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
    except xlrd.biffh.XLRDError as exc:
        if "encrypt" in str(exc).lower():
            raise UserError("protected_file") from exc
        raise UserError("corrupted_file") from exc
    selected = ctx.options.get("sheets")
    tables: list[RawTable] = []
    try:
        for s_idx in range(book.nsheets):
            sh = book.sheet_by_index(s_idx)
            if selected and sh.name not in selected:
                continue
            if sh.nrows > limits.max_rows + 1000:
                raise LimitExceeded("too_many_rows_sheet", sheet=sh.name, limit=limits.max_rows)
            if sh.ncols > limits.max_cols + 50:
                raise LimitExceeded("too_many_columns", limit=limits.max_cols, count=sh.ncols)
            grid: list[list[str | None]] = []
            srcs: list[str] = []
            for r in range(sh.nrows):
                out: list[str | None] = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            out.append(_encode(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                        except (ValueError, OverflowError):
                            out.append(None)
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        fmt = None
                        try:
                            xf = book.xf_list[cell.xf_index]
                            fmt = book.format_map[xf.format_key].format_str
                        except (IndexError, KeyError, AttributeError):
                            fmt = None
                        out.append(_encode(cell.value, fmt))
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        out.append(native_bool(bool(cell.value)))
                    elif cell.ctype in (xlrd.XL_CELL_TEXT,):
                        out.append(clean_text(cell.value) or None)
                    else:
                        out.append(None)
                grid.append(out)
                srcs.append(f"sheet={sh.name};row={r + 1}")
            if not grid:
                continue
            notes: list[dict] = []
            refs = [
                f"{get_column_letter(c0 + 1)}{r0 + 1}:{get_column_letter(c1)}{r1}" for (r0, r1, c0, c1) in sh.merged_cells
            ]
            _apply_merges(grid, refs, notes)
            width = max(len(r) for r in grid)
            col_refs = [get_column_letter(i + 1) for i in range(width)]
            for t in structure_grid(grid, srcs, f"{name} › {sh.name}", "sheet", col_refs=col_refs):
                t.notes.extend(notes)
                t.note("sheet", name=sh.name)
                t.note("xls_cached_results")
                tables.append(t)
            ctx.progress("extract", 0.2 + 0.6 * (s_idx + 1) / max(1, book.nsheets))
            book.unload_sheet(s_idx)
    finally:
        book.release_resources()
    if not tables:
        raise UserError("no_data_found")
    return tables
