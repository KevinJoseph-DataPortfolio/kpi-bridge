"""DOCX (tables + text), TXT (delimited or free text) and JSON (flattened lists of objects)."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from ..errors import LimitExceeded, UserError
from .base import IngestContext, RawTable, clean_text, native_bool, native_number
from .csvtab import looks_delimited, parse_delimited_text, read_text
from .tables import structure_grid
from .text import extract_text_values


def parse_docx(path: Path, name: str, ctx: IngestContext) -> list[RawTable]:
    import docx  # python-docx: lxml parser without entity resolution

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise UserError("corrupted_file") from exc
    tables: list[RawTable] = []
    for t_idx, table in enumerate(document.tables, start=1):
        grid: list[list[str | None]] = []
        srcs: list[str] = []
        for r_idx, row in enumerate(table.rows, start=1):
            cells: list[str | None] = []
            prev_tc = None
            for cell in row.cells:
                # horizontally merged cells: python-docx repeats the same cell
                if prev_tc is not None and cell._tc is prev_tc:
                    cells.append(None)
                    continue
                prev_tc = cell._tc
                cells.append(clean_text(cell.text) or None)
            if len(cells) > ctx.limits.max_cols:
                raise LimitExceeded("too_many_columns", limit=ctx.limits.max_cols, count=len(cells))
            grid.append(cells)
            srcs.append(f"table={t_idx};row={r_idx}")
        if len(grid) > ctx.limits.max_rows + 1000:
            raise LimitExceeded("too_many_rows", limit=ctx.limits.max_rows)
        tables.extend(structure_grid(grid, srcs, f"{name} · tableau {t_idx}", "docx"))
    if not tables:
        lines = [(p.text, f"paragraph={i}") for i, p in enumerate(document.paragraphs, start=1)]
        t = extract_text_values(lines, name)
        if t is not None:
            tables.append(t)
    if not tables:
        raise UserError("no_data_found_text")
    return tables


def parse_txt(path: Path, name: str, ctx: IngestContext) -> list[RawTable]:
    text, enc_note = read_text(path, ctx)
    if looks_delimited(text):
        return parse_delimited_text(text, name, ctx, ctx.options.get("delimiter"), enc_note)
    lines = [(ln, f"line={i}") for i, ln in enumerate(text.splitlines(), start=1)]
    t = extract_text_values(lines, name)
    if t is None:
        raise UserError("no_data_found_text")
    return [t]


def _depth(obj: Any, limit: int, level: int = 0) -> None:
    if level > limit:
        raise UserError("json_too_deep", limit=limit)
    if isinstance(obj, dict):
        for v in obj.values():
            _depth(v, limit, level + 1)
    elif isinstance(obj, list):
        for v in obj[:1000]:
            _depth(v, limit, level + 1)


def _flatten(obj: dict, prefix: str = "", out: dict | None = None, notes: set | None = None, level: int = 0) -> dict:
    out = {} if out is None else out
    for k, v in obj.items():
        key = f"{prefix}.{k}" if prefix else str(k)
        if isinstance(v, dict) and level < 3:
            _flatten(v, key, out, notes, level + 1)
        elif isinstance(v, list):
            if all(not isinstance(x, (dict, list)) for x in v):
                out[key] = json.dumps(v, ensure_ascii=False)
                if notes is not None:
                    notes.add(("list_as_text", key))
            elif notes is not None:
                notes.add(("nested_objects_unsupported", key))
        else:
            out[key] = v
    return out


def _encode(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return native_bool(v)
    if isinstance(v, (int, float)):
        return native_number(v) or None
    return clean_text(str(v)) or None


def parse_json(path: Path, name: str, ctx: IngestContext) -> list[RawTable]:
    raw = path.read_bytes()
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise UserError("json_invalid") from exc
    _depth(data, ctx.limits.max_json_depth)
    records: list | None = None
    chosen_key = None
    if isinstance(data, list):
        records = data
    elif isinstance(data, dict):
        candidates = {k: v for k, v in data.items() if isinstance(v, list) and v and all(isinstance(x, dict) for x in v)}
        wanted = ctx.options.get("json_key")
        if wanted and wanted in candidates:
            chosen_key = wanted
        elif candidates:
            chosen_key = max(candidates, key=lambda k: len(candidates[k]))
        if chosen_key:
            records = candidates[chosen_key]
    if not records or not all(isinstance(r, dict) for r in records):
        raise UserError("json_unsupported_structure")
    if len(records) > ctx.limits.max_rows:
        raise LimitExceeded("too_many_rows", limit=ctx.limits.max_rows)
    notes: set = set()
    flat = [_flatten(r, notes=notes) for r in records]
    columns: list[str] = []
    for r in flat:
        for k in r:
            if k not in columns:
                columns.append(k)
                if len(columns) > ctx.limits.max_cols:
                    raise LimitExceeded("too_many_columns", limit=ctx.limits.max_cols, count=len(columns))
    rows: list[list[str | None]] = [list(columns)]
    srcs = ["header"]
    for i, r in enumerate(flat):
        rows.append([_encode(r.get(c)) for c in columns])
        srcs.append(f"record={i}" + (f";key={chosen_key}" if chosen_key else ""))
    t = RawTable(name=f"{name}", kind="json", rows=rows, row_src=srcs, header_row=0, col_refs=columns)
    if chosen_key:
        t.note("json_key", key=chosen_key)
    for code, key in sorted(notes):
        t.note(code, field=key)
    return [t]
